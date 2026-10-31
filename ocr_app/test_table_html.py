"""合并单元格表格单测：HTML 解析 / 网格布局 / Word 合并单元格渲染 / 导出向后兼容。

验证 Task #39 的不变量：
- parse_html_tables 正确解析 rowspan / colspan；
- layout_table 把跨行/跨列单元格布局为规整网格与 anchor；
- _add_html_table 渲染出真实合并单元格（docx XML 含 gridSpan / vMerge）；
- md_to_docx 对纯文本 / Markdown 表格 / 内嵌 HTML 表格均不退化。
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))  # ocr_app

from utils.table_html import parse_html_tables, layout_table
from utils.export_docx import md_to_docx, _add_html_table
from docx import Document


HTML_ROWSPAN = """
<table>
<tr><td rowspan="2">A</td><td>B</td></tr>
<tr><td>C</td></tr>
</table>
"""

HTML_COLSPAN = """
<table>
<tr><td colspan="2">H</td></tr>
<tr><td>X</td><td>Y</td></tr>
</table>
"""


def test_parse_html_tables():
    tables = parse_html_tables(HTML_ROWSPAN)
    assert len(tables) == 1
    rows = tables[0]
    assert len(rows) == 2
    assert rows[0][0]["rowspan"] == 2
    assert rows[0][0]["text"] == "A"
    assert rows[0][1]["text"] == "B"
    assert rows[1][0]["text"] == "C"
    # 文本空白被规整
    assert parse_html_tables("<table><tr><td>  hello   world </td></tr></table>")[0][0][0]["text"] == "hello world"


def test_layout_rowspan():
    raw = parse_html_tables(HTML_ROWSPAN)[0]
    grid, nrows, ncols, anchors = layout_table(raw)
    assert nrows == 2 and ncols == 2
    # anchor A 跨 2 行 1 列
    a = [a for a in anchors if a[4] == "A"][0]
    assert a[2] == 2 and a[3] == 1  # rowspan, colspan
    # (1,0) 被 A 覆盖
    assert grid[1][0] == grid[0][0]


def test_layout_colspan():
    raw = parse_html_tables(HTML_COLSPAN)[0]
    grid, nrows, ncols, anchors = layout_table(raw)
    assert nrows == 2 and ncols == 2
    h = [a for a in anchors if a[4] == "H"][0]
    assert h[3] == 2  # colspan=2
    assert grid[0][0] == grid[0][1]  # H 覆盖两列


def test_render_html_table_merges():
    doc = Document()
    _add_html_table(doc, HTML_ROWSPAN)
    xml = doc.element.xml
    assert "gridSpan" in xml or "vMerge" in xml, "应出现合并单元格标记"
    assert "vMerge" in xml, "rowspan 应产生纵向合并 vMerge"


def test_render_html_colspan_merges():
    doc = Document()
    _add_html_table(doc, HTML_COLSPAN)
    xml = doc.element.xml
    assert "gridSpan" in xml, "colspan 应产生横向合并 gridSpan"


def test_md_to_docx_backward_compat():
    import tempfile
    d = tempfile.mkdtemp()
    p = os.path.join(d, "out.docx")
    md_to_docx("# 标题\n\n一段中文正文。\n\n| a | b |\n|---|---|\n| 1 | 2 |\n", p)
    assert os.path.exists(p)
    doc = Document(p)
    txt = "\n".join(par.text for par in doc.paragraphs)
    assert "标题" in txt and "中文正文" in txt
    # 内嵌 HTML 表格也能渲染而不报错
    p2 = os.path.join(d, "out2.docx")
    md_to_docx("前面一段。\n\n" + HTML_ROWSPAN + "\n\n后面一段。", p2)
    doc2 = Document(p2)
    assert "gridSpan" in doc2.element.xml or "vMerge" in doc2.element.xml


def main():
    test_parse_html_tables()
    test_layout_rowspan()
    test_layout_colspan()
    test_render_html_table_merges()
    test_render_html_colspan_merges()
    test_md_to_docx_backward_compat()
    print("ALL PASS: test_table_html")


if __name__ == "__main__":
    main()
