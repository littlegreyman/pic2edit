"""导出：把 Markdown 结果保存为 .md / .txt，以及轻量转换为 .docx。

Markdown -> docx 转换器覆盖常见元素：
- #/##/### 标题
- - / * 无序列表
- 1. 有序列表
- | a | b | 表格
- **加粗** 行内格式
"""
import os
import re
from typing import List

from docx import Document
from docx.shared import Pt

from .table_html import parse_html_tables, layout_table


def save_text(text: str, out_path: str):
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(text)


def _add_runs(paragraph, line: str):
    """处理 **加粗** 行内格式，逐段写入。"""
    parts = re.split(r"(\*\*.*?\*\*)", line)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def md_to_docx(text: str, out_path: str):
    doc = Document()
    _populate_doc(doc, text)
    doc.save(out_path)


def md_pages_to_docx(page_texts: List[str], out_path: str, page_break: bool = True):
    """把多页 Markdown 文本装配成一个 Word 文档，页与页之间插入真实分页符。

    page_texts: 各页的 Markdown 文本列表（顺序即页序）。
    第 1 页之前不插分页；其余页之前若 page_break=True 则 add_page_break()。
    """
    doc = Document()
    for i, text in enumerate(page_texts):
        if i > 0 and page_break:
            doc.add_page_break()
        _populate_doc(doc, text or "")
    doc.save(out_path)


def _split_html_tables(text: str):
    """把文本切成片段：("md", 段) 或 ("html", <table>...</table>)。

    用于在同一篇结果里同时支持 Markdown 与 VLM 输出的 HTML 表格（含合并单元格）。
    """
    out = []
    pos = 0
    for m in re.finditer(r"<table\b[^>]*>.*?</table>", text, re.DOTALL | re.IGNORECASE):
        if m.start() > pos:
            out.append(("md", text[pos:m.start()]))
        out.append(("html", m.group(0)))
        pos = m.end()
    if pos < len(text):
        out.append(("md", text[pos:]))
    return out


def _add_html_table(doc, html: str):
    """把一段 HTML（可含多个 <table>）渲染为 Word 表格，支持合并单元格。"""
    for raw in parse_html_tables(html):
        if not raw:
            continue
        grid, nrows, ncols, anchors = layout_table(raw)
        if nrows == 0 or ncols == 0:
            continue
        try:
            table = doc.add_table(rows=nrows, cols=ncols)
        except Exception:
            continue
        table.style = "Table Grid"
        for (r, c, rs, cs, text) in anchors:
            try:
                cell = table.cell(r, c)
                if rs > 1 or cs > 1:
                    cell.merge(table.cell(r + rs - 1, c + cs - 1))
                cell.text = text or ""
            except Exception:
                pass
        doc.add_paragraph()


def _populate_doc(doc, text: str):
    """按片段填充 doc：Markdown 段走原解析；HTML <table> 段走 _add_html_table。"""
    for kind, content in _split_html_tables(text):
        if kind == "html":
            _add_html_table(doc, content)
        else:
            _parse_md_segment(doc, content)


def _parse_md_segment(doc, text: str):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        # 表格：连续以 | 开头/分隔的行
        if line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            # 去掉分隔行（---）
            rows = [r for r in table_lines if not re.fullmatch(r"\|\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|", r)]
            if rows:
                cells = [c.strip() for c in rows[0].strip("|").split("|")]
                ncol = len(cells)
                table = doc.add_table(rows=1, cols=ncol)
                table.style = "Table Grid"
                hdr = table.rows[0].cells
                for c, val in enumerate(cells):
                    hdr[c].text = val
                for row in rows[1:]:
                    vals = [c.strip() for c in row.strip("|").split("|")]
                    vals += [""] * (ncol - len(vals))
                    cells_row = table.add_row().cells
                    for c, val in enumerate(vals[:ncol]):
                        cells_row[c].text = val
                doc.add_paragraph()
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            p = doc.add_heading(level=min(level, 6))
            _add_runs(p, m.group(2))
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, m.group(1))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^\d+\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, m.group(1))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        _add_runs(p, line)
        i += 1
